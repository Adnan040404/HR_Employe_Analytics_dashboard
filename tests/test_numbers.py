"""Tests for the HR attrition project.

The expected values are worked out separately (plain csv module, hand-checked
arithmetic) instead of calling the same functions the script uses, and the Excel
workbook's own cached results are compared with the Python numbers.
"""

import csv
import os
import sys
import warnings

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import verify_numbers as vn  # noqa: E402


@pytest.fixture(scope="module")
def df():
    return vn.load()


@pytest.fixture(scope="module")
def raw():
    with open(os.path.join(ROOT, "data", "hr_attrition.csv"), newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def workbook():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about slicers it can't read
        return openpyxl.load_workbook(os.path.join(ROOT, "HR-Employee-Analysis.xlsx"),
                                      read_only=True, data_only=True)


# ------------------------------------------------------------------- the data
def test_headline_numbers(raw):
    assert len(raw) == 1470
    assert sum(r["Attrition"] == "Yes" for r in raw) == 237
    assert len({r["Employee Number"] for r in raw}) == 1470          # no duplicate employees
    assert {r["Attrition"] for r in raw} == {"Yes", "No"}


def test_department_rates_by_hand(raw):
    counts = {}
    for r in raw:
        n, l = counts.get(r["Department"], (0, 0))
        counts[r["Department"]] = (n + 1, l + (r["Attrition"] == "Yes"))
    assert counts == {"Sales": (446, 92), "Research & Development": (961, 133), "Human Resources": (63, 12)} \
        or counts == {"Sales": (446, 92), "R&D": (961, 133), "HR": (63, 12)}


def test_every_group_table_adds_up(df):
    for column in ["Department", "Age band", "Marital Status", "Job Role", "Education", "Over Time"]:
        t = vn.rate_table(df, column)
        assert t["employees"].sum() == 1470, column
        assert t["leavers"].sum() == 237, column
        assert t["share_of_leavers"].sum() == pytest.approx(1.0)


def test_age_band_matches_workbook_boundaries(df):
    assert df.loc[df["Age"] == 24, "Age band"].eq("Under 25").all()
    assert df.loc[df["Age"] == 25, "Age band"].eq("25 - 34").all()
    assert df.loc[df["Age"] == 54, "Age band"].eq("45 - 54").all()
    assert df.loc[df["Age"] == 55, "Age band"].eq("Over 55").all()


# ------------------------------------------- the point of the project (the bug)
def test_headcount_is_not_attrition(df):
    """The first dashboard's Age-Group, Job Role and Marital Status pivots counted people.
    Counting rows always totals 1470; the number that left totals 237."""
    for column in ["Age band", "Job Role", "Marital Status"]:
        t = vn.rate_table(df, column)
        assert t["employees"].sum() == 1470
        assert t["leavers"].sum() == 237


def test_biggest_group_is_not_the_highest_rate(df):
    dept = vn.rate_table(df, "Department")
    biggest_share = dept["share_of_leavers"].idxmax()
    highest_rate = dept["rate"].idxmax()
    assert biggest_share != highest_rate          # R&D has most leavers, Sales has the highest rate
    assert dept.loc[highest_rate, "rate"] == pytest.approx(92 / 446)

    age = vn.rate_table(df, "Age band")
    assert age["rate"].idxmax() == "Under 25"      # 39%, though "25 - 34" has the most leavers
    assert age["leavers"].idxmax() == "25 - 34"


# ------------------------------------------------------------------ the stats
def test_wilson_interval_known_values():
    lo, hi = vn.wilson(5, 10)
    assert lo == pytest.approx(0.2366, abs=1e-3) and hi == pytest.approx(0.7634, abs=1e-3)
    lo, hi = vn.wilson(0, 20)
    assert lo == pytest.approx(0.0, abs=1e-9) and 0.15 < hi < 0.20
    assert vn.wilson(0, 0) == (0.0, 0.0)
    for leavers, n in [(1, 3), (30, 100), (237, 1470)]:
        lo, hi = vn.wilson(leavers, n)
        assert 0 <= lo <= leavers / n <= hi <= 1


def test_drivers_ignore_small_groups_and_are_sorted(df):
    d = vn.drivers(df, vn.FEATURES)
    assert (d["employees"] >= vn.MIN_GROUP).all()
    assert d["lift"].is_monotonic_decreasing
    top = d.iloc[0]
    assert (top["factor"], top["value"]) == ("Job Role", "Sales Representative")
    assert top["lift"] == pytest.approx((33 / 83) / (237 / 1470))


def test_risk_flags(df):
    rp = vn.risk_profile(df)
    assert list(rp.index) == [0, 1, 2, 3]
    assert rp["employees"].sum() == 1470 and rp["leavers"].sum() == 237
    assert rp["rate"].is_monotonic_increasing      # more flags, higher attrition
    assert rp.loc[0, "rate"] < 0.08 and rp.loc[3, "rate"] > 0.6
    assert rp.loc[3, "employees"] == 40            # the top group is small: say so in the README


# -------------------------------------------- the Excel workbook agrees with Python
def test_workbook_rates_sheet_matches_python(workbook, df):
    ws = workbook["Attrition Rates"]
    rows = [r[:5] for r in ws.iter_rows(min_row=1, values_only=True) if r[0] is not None]
    lookup = {}
    block = None
    for a, b, c, d, e in rows:
        if b == "Employees":
            block = a
        elif a not in ("Total",) and isinstance(b, (int, float)):
            lookup[(block, a)] = (b, c, d)
    for label, column in [("Department", "Department"), ("Age band", "Age band"),
                          ("Marital status", "Marital Status"), ("Job role", "Job Role"),
                          ("Education", "Education"), ("Overtime", "Over Time")]:
        table = vn.rate_table(df, column)
        for group, row in table.iterrows():
            key_group = {"Sales": "Sales", "Research & Development": "R&D", "Human Resources": "HR"}.get(group, group) \
                if column == "Department" else group
            employees, leavers, rate = lookup[(label, key_group)]
            assert (employees, leavers) == (row["employees"], row["leavers"]), (label, group)
            assert rate == pytest.approx(row["rate"])


def test_workbook_pivots_now_count_leavers(workbook):
    for name in ["Attrition By Job Role", "Attrition By Age-Group", "Attrition By Marriage Status"]:
        rows = [r for r in workbook[name].iter_rows(values_only=True) if r[0] is not None]
        assert rows[-1][0] == "Grand Total" and rows[-1][1] == 237, name   # was 1470 (headcount)
